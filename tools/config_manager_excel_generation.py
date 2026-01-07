import json
import os
import argparse
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side

skip_fields = {
    "root": ["version"],
    "sensor_config": ["version", "source", "sensor_i2c_bus", "sensor_i2c_addr", "af_i2c_bus", "af_i2c_addr"],
    "application_settings": ["version", "debug"],
    "iq_settings": ["version", "Automatic_algorithms"],
    "osd": ["version"],
    "encoding": ["version"],
    "privacy_mask": ["version"]
}

# Define UI component types
TYPE_MAPPING = {
    str: "text box",
    bool: "checkbox",
    int: "slider",
    float: "slider"
}

# Fill styles for Excel
HEADER_FILL = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
SECTION_FILLS = [
    PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
    PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
]

THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# User-defined additional dependency rules
EXTERNAL_DEPENDENCY_RULES = [
    {
        "type": "max_array_length",
        "section": "encoded_output_streams",
        "limit": 6,
        "message": "No more than 6 output streams allowed"
    },
    {
        "type": "max_array_length",
        "section": "application_input_streams",
        "limit": 6,
        "message": "No more than 6 application input streams allowed"
    },       
    {
        "type": "mutual_exclusive",
        "params": ["application_settings.dis.enabled", "application_settings.eis.enabled"],
        "message": "DIS and EIS cannot be enabled at the same time"
    }
]

# Utility to infer parameter type

def infer_type(value):
    if isinstance(value, str):
        if value.endswith(".json"):
            return "File loader"
        return TYPE_MAPPING.get(str, "text box")
    return TYPE_MAPPING.get(type(value), "text box")


def apply_external_dependencies(full_param_path, is_enable_field=False):
    messages = []
    print(f"DEBUG: apply_external_dependencies called for {full_param_path}, is_enable_field={is_enable_field}")
    for rule in EXTERNAL_DEPENDENCY_RULES:
        if rule["type"] == "max_array_length":
            if rule["section"] in full_param_path:
                messages.append(rule["message"])
        elif rule["type"] == "mutual_exclusive" and is_enable_field:
            print(f"DEBUG: Checking mutual exclusion for {full_param_path} in {rule['params']}")
            if full_param_path in rule["params"]:
                print(f"DEBUG: Match found for mutual exclusion on {full_param_path}")
                messages.append(rule["message"])
    return "; ".join(messages) if messages else None


def extract_parameters(section_name, data, skip_fields=None, module="top", parent_enable_path=None):
    rows = []
    skip_fields = skip_fields or []

    if isinstance(data, dict):
        local_enable_field = None
        for k in ("enable", "enabled"):
            if k in data:
                local_enable_field = f"{module}.{k}"
                break

        for key, value in data.items():
            if key in skip_fields:
                continue

            full_param_path = f"{module}.{key}"
            print(f"DEBUG: Processing '{full_param_path}' - is_enable_field: {local_enable_field}")
            is_enable_field = key in ("enable", "enabled")

            # Determine dependency string
            dependency_msg = None
            if local_enable_field and not is_enable_field:
                dependency_msg = "available only when enable is selected"

            # Apply external rules
            external_msg = apply_external_dependencies(full_param_path, is_enable_field=is_enable_field)
            print(f"DEBUG: Dependencies for '{full_param_path}': {external_msg}")

            # Combine
            combined_dependency = "; ".join(filter(None, [dependency_msg, external_msg])) or None

            if isinstance(value, dict):
                rows.append({
                    "Section": section_name,
                    "Module": module,
                    "Parameter": key,
                    "type": "sub-section",
                    "dependency": None
                })
                rows.extend(extract_parameters(section_name, value, skip_fields, module=f"{module}.{key}", parent_enable_path=local_enable_field))

            elif isinstance(value, list) and value and isinstance(value[0], dict):
                rows.append({
                    "Section": section_name,
                    "Module": module,
                    "Parameter": key,
                    "type": "addable sub-section",
                    "dependency": None
                })
                rows.extend(extract_parameters(section_name, value[0], skip_fields, module=f"{module}.{key}", parent_enable_path=local_enable_field))


            else:
                rows.append({
                    "Section": section_name,
                    "Module": module,
                    "Parameter": key,
                    "type": infer_type(value),
                    "dependency": combined_dependency
                })

    return rows


def write_to_excel(rows, output_excel):
    wb = Workbook()
    ws = wb.active
    ws.title = "Configuration"

    headers = ["Section", "Module", "Parameter", "type", "dependency", "file"]
    ws.append(headers)

    for col_num, col in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.border = THIN_BORDER

    section_colors = defaultdict(lambda: SECTION_FILLS[0])
    color_index = 0
    current_section = None

    for row_data in rows:
        if row_data["Section"] != current_section:
            current_section = row_data["Section"]
            color_index = (color_index + 1) % len(SECTION_FILLS)
        fill = SECTION_FILLS[color_index]

        row = [row_data.get(header, "") for header in headers]
        ws.append(row)
        row_idx = ws.max_row

        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            cell.border = THIN_BORDER

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max_length + 2

    wb.save(output_excel)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_json", help="Path to the main JSON configuration file")
    parser.add_argument("output_excel", help="Path to the output Excel file")
    args = parser.parse_args()

    base_dir = os.path.dirname(args.input_json)

    with open(args.input_json, "r") as f:
        main_data = json.load(f)

    all_rows = []

    # First gather top-level summary sections
    for key, path in main_data.items():
        if key in skip_fields.get("root", []):
            continue
        if key == "encoded_output_streams":
            all_rows.append({
                "Section": "top",
                "Module": "top",
                "Parameter": "encoded output streams",
                "type": "section",
                "dependency": None,
                "file": "profile.json"
            })

            all_rows.append({
                "Section": "encoded output streams",
                "Module": "top",
                "Parameter": "output_stream",
                "type": "addable sub-section",
                "dependency": None,
                "file": "profile.json"
            })

            first_stream = main_data["encoded_output_streams"][0]
            for subkey, subpath in first_stream.items():
                if subkey == "stream_id":
                    continue
                filename = os.path.basename(subpath)
                full_path = os.path.join(base_dir, filename)
                if os.path.exists(full_path) and os.path.isfile(full_path):
                    with open(full_path) as f:
                        data = json.load(f)
                    all_rows.append({
                        "Section": "output_stream",
                        "Module": "top",
                        "Parameter": subkey,
                        "type": "sub-section",
                        "dependency": None,
                        "file": filename
                    })
                    all_rows.extend(extract_parameters(subkey, data, skip_fields.get(key, []), module=key))


        elif isinstance(path, str) and path.endswith(".json"):
            filename = os.path.basename(path)
            full_path = os.path.join(base_dir, filename)
            if os.path.exists(full_path):
                all_rows.append({
                    "Section": "top",
                    "Module": "top",
                    "Parameter": key,
                    "type": "section",
                    "dependency": None,
                    "file": filename
                    
                })
                with open(full_path) as f:
                    data = json.load(f)
                all_rows.extend(extract_parameters(key, data, skip_fields.get(key, []), module=key))


    write_to_excel(all_rows, args.output_excel)


if __name__ == "__main__":
    main()
